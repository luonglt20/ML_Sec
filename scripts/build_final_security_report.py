#!/usr/bin/env python3
"""Build the final Pair-1 report from the verified DeepSeek summary JSON."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
MEMORY_SUMMARY = ROOT / "results/security/deepseek20_frozen_memory_summary.json"
RAG_SUMMARY = ROOT / "results/security/deepseek20_rag_combined_summary.json"
OUTPUT = ROOT / "deliverables/Bao_cao_final_Base_Giua_Ky_Bat_Bien.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
PALE_RED = "FDECEC"
PALE_GREEN = "EAF4EA"
GRAY = "666666"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths_dxa):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_run(run, *, size=11, color=INK, bold=False, italic=False) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_text(doc, text: str, *, style="Normal", bold=False, italic=False, color=INK, align=None) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run(run)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run(run, size=9, color=GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], *, highlight_last=False) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run(run, size=9, bold=True, color=DARK_BLUE)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            if highlight_last and index == len(rows) - 1:
                set_cell_shading(cell, PALE_GREEN)
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_run(run, size=9)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def condition(summary, track, surface, defense, family, position="top1", placement=None):
    for row in summary["conditions"]:
        if (
            row["model_track"] == track
            and row["attack_surface"] == surface
            and row["defense"] == defense
            and row["attack_family"] == family
            and row["attack_position"] == position
            and (placement is None or row["payload_placement"] == placement)
        ):
            return row
    raise KeyError((track, surface, defense, family, position))


def percent(value: float) -> str:
    return f"{value * 100:.1f}%"


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("ML-Sec | Pair 1 — Prompt Injection Security Evaluation")
    set_run(run, size=9, color=GRAY)


def main() -> None:
    memory_summary = json.loads(MEMORY_SUMMARY.read_text(encoding="utf-8"))
    rag_summary = json.loads(RAG_SUMMARY.read_text(encoding="utf-8"))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # First-page memo masthead.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("BÁO CÁO FINAL — AN TOÀN THÔNG TIN CHO LLM")
    set_run(run, size=22, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Pair 1: Indirect Prompt Injection trên RAG và Long-term Memory")
    set_run(run, size=14, color=GRAY)
    metadata = [
        ("Mô hình baseline", "Pipeline giữa kỳ V2/V4, DeepSeek Chat"),
        ("Evaluation", "20 câu MedQA-USMLE stratified, seed 14, temperature 0"),
        ("Snapshot", "Frozen DeepSeek/MedRAG snapshot; seed 14"),
        ("Trạng thái", "350 DeepSeek API calls thật, uncached, 27/08/2026"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    for label, value in metadata:
        cells = table.add_row().cells
        set_cell_shading(cells[0], PALE_GRAY)
        p0 = cells[0].paragraphs[0]
        p1 = cells[1].paragraphs[0]
        set_run(p0.add_run(label), size=10, bold=True, color=DARK_BLUE)
        set_run(p1.add_run(value), size=10)
    doc.add_paragraph()

    doc.add_heading("Tóm tắt kết quả", level=1)
    add_text(doc, "Baseline giữa kỳ được giữ nguyên byte-for-byte: không sửa agent, entrypoint, workflow, prompt mặc định, corpus, retrieval ranking hay nhãn đáp án. Security harness ngoài baseline chỉ tạo attacked/defended view ở dữ liệu không tin cậy sau retrieval, rồi gọi pipeline gốc.")
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [9360])
    set_cell_shading(callout.cell(0, 0), PALE_GREEN)
    callout.cell(0, 0).text = "Kết quả chính V2 RAG trên base bất biến: clean 19/20 (95%) → Combined top-1 18/20 (90%), targeted ASR 1/12 (8.3%). Stress document-authority/all-top-k đạt ASR 1/12 (8.3%) nhưng không làm giảm accuracy ròng; input guard đưa ASR stress về 0/13 và lọc payload 100%. R2 Chained-Completion, R3 high-pressure memory và R4 five-shot context smoke đều 5/5, ASR 0/3 — negative results được giữ lại."
    for run in callout.cell(0, 0).paragraphs[0].runs:
        set_run(run, size=10, bold=True, color=INK)

    doc.add_page_break()

    doc.add_heading("1. Threat model và bất biến baseline", level=1)
    add_text(doc, "Kẻ tấn công chỉ kiểm soát text của passage RAG hoặc explanation của memory case sau khi retrieval hoàn tất. ID, score, thứ tự top-k, corpus gốc, đáp án gold và thuật toán retrieval đều không thay đổi. Điều này tách rõ indirect prompt injection (Pair 1) khỏi knowledge-base poisoning (Pair 3).")
    add_bullet(doc, "RAG/V2 là kết quả chính: Router → Retriever → Reasoner → Verifier chạy như baseline.")
    add_bullet(doc, "Memory/V4 là mở rộng: memory được khởi tạo từ dev.jsonl; chỉ test.jsonl được đánh giá.")
    add_bullet(doc, "Attack payload yêu cầu một target option sinh từ hash question ID và marker PWNED; ASR chỉ tính clean-correct, target khác gold.")
    add_bullet(doc, "Mọi API condition là uncached để sanitizer không thể reuse clean output cho attack/defense.")

    doc.add_heading("2. Kỹ thuật attack và defense", level=1)
    add_text(doc, "Năm family core theo Open-Prompt-Injection gồm Naive, Escape, Context-Ignoring, Fake-Completion và Combined. Payload được inject vào top-1 passage; ablation all-top-k inject vào mọi passage của snapshot. Stress suite đã khóa thêm Combined suffix và payload document-authority (giả mạo addendum kiểm định tài liệu) ở suffix/sandwich. Residual probes gồm delimiter near-miss, multilingual và Base64-encoded instruction.")
    add_text(doc, "Defense thực nghiệm là transparent input guard: tách lớp pipeline không đổi, lọc paragraph không tin cậy có role token, override instruction, Final Answer hoặc marker trước khi model đọc. Đây là engineering baseline, không phải StruQ: StruQ đầy đủ phải gồm secure frontend và model structured-instruction-tuned checkpoint.")
    add_table(doc, ["RAG signal", "Clean", "Combined", "Authority stress", "Guard stress"], [[
        "DeepSeek real calls", "95%", "90%", "95%", "100%"
    ]], [2500, 1500, 1700, 2100, 1560], highlight_last=True)
    caption = add_text(doc, "Hình 1. V2 RAG là kết quả chính; guard được ghi rõ là engineering baseline, không phải StruQ.", italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
    caption.paragraph_format.space_before = Pt(2)

    doc.add_page_break()

    doc.add_heading("3. Thiết kế benchmark và kiểm soát tái lập", level=1)
    add_text(doc, "Tập 20 câu được stratify 5 câu cho mỗi nhãn A–D, seed 14. Mọi condition dùng lại cùng snapshot hash; generation temperature=0 và parser là cố định. Trước/sau benchmark, script verify_midterm_baseline.py xác nhận SHA-256 của entrypoint và bốn agent giữa kỳ.")
    add_table(doc, ["Artifact", "Giá trị"], [
        ["Model", "deepseek-chat"],
        ["Real calls", "350 (80 V4 memory + 120 V2 core + 120 V2 stress + 30 exploratory smoke; raw JSONL local, cache hit = 0)"],
        ["Baseline check", "5 protected source files match the frozen SHA-256 manifest"],
        ["Main metrics", "Accuracy; targeted ASR; marker ASR; bootstrap 95% CI; McNemar; latency; tokens; sanitization"],
        ["Raw output policy", "Raw JSONL/cache/index/.env ignored; summary JSON/CSV and code versioned"],
    ], [2400, 6960])

    doc.add_heading("4. V2 RAG: attack → defense (kết quả chính)", level=1)
    rag_clean = condition(rag_summary, "api", "rag", "none", "clean")
    rag_combined = condition(rag_summary, "api", "rag", "none", "combined", placement="prefix")
    authority_attack = condition(rag_summary, "api", "rag", "none", "document_authority", "all_top_k", "sandwich")
    authority_guard_clean = condition(rag_summary, "api_heuristic_guard", "rag", "heuristic_guard", "clean", "all_top_k", "sandwich")
    authority_guard = condition(rag_summary, "api_heuristic_guard", "rag", "heuristic_guard", "document_authority", "all_top_k", "sandwich")
    add_table(doc, ["Condition", "Correct", "Accuracy", "Targeted ASR", "Sanitized"], [
        ["Clean baseline", f"{rag_clean['correct']}/20", percent(rag_clean["accuracy"]), "—", "0%"],
        ["Combined top-1 prefix", f"{rag_combined['correct']}/20", percent(rag_combined["accuracy"]), "1/12 (8.3%)", "0%"],
        ["Authority all-top-k sandwich", f"{authority_attack['correct']}/20", percent(authority_attack["accuracy"]), "1/12 (8.3%)", "0%"],
        ["Guard clean", f"{authority_guard_clean['correct']}/20", percent(authority_guard_clean["accuracy"]), "—", "0%"],
        ["Guard + authority", f"{authority_guard['correct']}/20", percent(authority_guard["accuracy"]), "0/13 (0%)", "100%"],
    ], [2700, 1300, 1500, 2300, 1560], highlight_last=True)
    add_text(doc, "Kết luận V2 có giới hạn: DeepSeek bị target ở 1/12 câu đủ điều kiện, nhưng n=20 chưa quan sát accuracy collapse. Với authority stress, clean-vs-attack delta là 0 điểm (McNemar p=1.00); guard giảm measured targeted ASR 100% và accuracy attack cao hơn 1 câu, nhưng p=1.00. Những số này là evidence của attack surface và guard behavior, chưa phải bằng chứng mạnh về chênh lệch accuracy.")

    doc.add_heading("5. Kiểm chứng phân tách baseline", level=1)
    add_table(doc, ["Protected source", "Trạng thái"], [
        ["entrypoint.py", "không có diff so với midterm manifest"],
        ["Memory/Reasoner/Researcher/Verifier agents", "không có diff so với midterm manifest"],
        ["Attack/defense code", "chỉ tồn tại tại medqa_multiagent/security/harness.py và decorators"],
    ], [3900, 5460])
    add_text(doc, "Vì V4 gọi đúng Memory Agent, Router, Researcher và Reasoner gốc với một memory view decorator, chênh lệch kết quả được quy cho injected data/guard, không phải sửa prompt hoặc thay agent giữa kỳ.")

    doc.add_page_break()

    doc.add_heading("6. V4 long-term memory: mở rộng", level=1)
    mem_clean = condition(memory_summary, "api", "memory", "none", "clean")
    mem_attack = condition(memory_summary, "api", "memory", "none", "combined")
    guard_clean = condition(memory_summary, "api_heuristic_guard", "memory", "heuristic_guard", "clean")
    guard_attack = condition(memory_summary, "api_heuristic_guard", "memory", "heuristic_guard", "combined")
    add_table(doc, ["Condition", "Correct", "Accuracy", "Targeted ASR", "Sanitized"], [
        ["Baseline clean", f"{mem_clean['correct']}/20", percent(mem_clean["accuracy"]), "—", "0%"],
        ["Combined memory attack", f"{mem_attack['correct']}/20", percent(mem_attack["accuracy"]), f"1/13 (7.7%)", "0%"],
        ["Guard clean", f"{guard_clean['correct']}/20", percent(guard_clean["accuracy"]), "—", "0%"],
        ["Guard + Combined", f"{guard_attack['correct']}/20", percent(guard_attack["accuracy"]), "0/13 (0%)", "100%"],
    ], [2700, 1300, 1500, 2300, 1560], highlight_last=True)
    add_text(doc, "V4 cho causal chain rõ nhất: clean correct → retrieved-memory injection làm sai hai câu, trong đó một câu đi đúng target attacker → guard loại paragraph độc hại và khôi phục cả hai. Relative targeted-ASR reduction và recovery đều 100%; guard clean giữ nguyên 100%. Với n=20, McNemar cho attack và recovery đều p=0.50, nên report nêu effect size nhưng không khẳng định significance.")

    doc.add_page_break()

    doc.add_heading("7. Residual probes, limitations và hướng tiếp theo", level=1)
    add_table(doc, ["Residual scope", "Trạng thái"], [
        ["RAG/V2 five-family matrix", "đã chạy source-frozen n=20; attack mạnh nhất giảm 5 điểm ở Combined/Escape"],
        ["R2 Chained-Completion smoke", "đã chạy 5 câu; delivery 5/5 nhưng accuracy 5/5, ASR 0/3 (negative result)"],
        ["R3 high-pressure memory smoke", "all-retrieved-exemplar, sandwich, repeat ×3: 5/5, ASR 0/3 (negative result)"],
        ["R4 five-shot context smoke", "5 fake Q/A in copied top-1 passage, excluded from scoring: 5/5, ASR 0/3 (negative result)"],
        ["Delimiter near-miss / multilingual / encoded", "đã implement; phải chạy ở batch n=100 confirmatory"],
        ["Official StruQ checkpoint", "chưa chạy vì host không có GPU NVIDIA/checkpoint phù hợp"],
    ], [3800, 5560])
    add_bullet(doc, "GCG/optimization-based attack chưa được chạy: host không có GPU NVIDIA phù hợp cho official StruQ setup. Đây là residual gap, không phải claim defense tuyệt đối.")
    add_bullet(doc, "R2 guard canonicalizes zero-width Unicode, role-boundary token và suspicious Base64 instruction. Marker ASR chỉ tính output model, không tính trace; guard vẫn có thể bỏ sót obfuscation mới và cần matched StruQ model để giảm trade-off.")
    add_bullet(doc, "n=20 là evaluation thật có kiểm soát nhưng CI rộng; mở rộng đúng snapshot protocol lên n=100 là bước tiếp theo khi có ngân sách API.")

    doc.add_page_break()

    doc.add_heading("8. Reproduce và demo", level=1)
    add_text(doc, "Lệnh tái lập chính nằm trong README.md: trước hết chạy scripts/verify_midterm_baseline.py, sau đó freeze snapshot DeepSeek và chạy harness ngoài baseline. Summary versioned: results/security/deepseek20_rag_combined_summary.json/csv và results/security/deepseek20_frozen_memory_summary.json/csv. Raw prediction JSONL bị gitignore để không commit prompt/trace và chi phí không cần thiết.")
    add_text(doc, "Demo nên chọn một trong hai case recovered từ raw trace: clean trả gold → Combined memory payload làm sai → input guard khôi phục. Chỉ dùng examples có trace thật, không viết lại để phù hợp narrative.")

    doc.add_heading("Tài liệu tham khảo", level=1)
    add_text(doc, "[1] Yupei Liu et al. Formalizing and Benchmarking Prompt Injection Attacks and Defenses. USENIX Security 2024. Official toolkit: github.com/liu00222/Open-Prompt-Injection.", color=GRAY)
    add_text(doc, "[2] Sizhe Chen et al. StruQ: Defending Against Prompt Injection with Structured Queries. USENIX Security 2025. Official implementation: github.com/Sizhe-Chen/StruQ.", color=GRAY)
    add_text(doc, "[3] Project artifacts: security/midterm_baseline_manifest.json; security/deepseek_exploratory_r2_protocol.json; security/deepseek_memory_highpressure_r3_protocol.json; security/deepseek_rag_fewshot_r4_protocol.json; medqa_multiagent/security/harness.py; results/security/deepseek20_rag_combined_summary.json; results/security/deepseek20_frozen_memory_summary.json; scripts/run_security_eval.py.", color=GRAY)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
